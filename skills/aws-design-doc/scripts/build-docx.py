#!/usr/bin/env python3
# Copyright Amazon.com, Inc. or its affiliates. All Rights Reserved.
# SPDX-License-Identifier: MIT-0
"""Build a Word version of a design document carrying the same AWS theme as the HTML build.

    python3 build-docx.py MY-DESIGN.md [--strict]

Pandoc converts the markdown (it handles tables, lists and code far better than anything
hand-rolled); python-docx then restyles the result, because pandoc has no notion of a palette.

WHY THE DIAGRAMS ARE RASTERISED. The HTML embeds the architecture SVGs inline, which is what
makes that file self-contained. Word cannot take an inline SVG from a converted document, so
each `<!-- aws-svg: ... -->` marker is rewritten to an image reference and the SVG is rendered
to a 3x PNG with the bundled converter first. 3x rather than 1x so it stays crisp when zoomed or printed.
`<!-- aws-svg: path.svg | Caption -->` adds a "Figure N. Caption" paragraph under the image and
sets the caption as the picture's alt text. Mermaid fences have no Word rendering: the source is
kept as a code block and the build says so.

The layout is a classic consulting design document on A4: a cover page (title, subtitle and a
Prepared by / Prepared for / Version / Status / Date / Classification table), a Document history
page holding the first table under `## Document history`, a Table of Contents page (a Word TOC
field, levels 1-2), then the body. `##` sections are Heading 1 and `###` Heading 2. Every page
after the cover carries a header (prepared for and by) and a footer (title, version,
classification, page number). Word fills the contents when it opens the file (it asks first).
Kinds `brief` and `refarch` are one page: a compact title block, then the body, no cover,
history or contents page. Optional YAML front matter (title, kind, subtitle, prepared_by,
prepared_for, version, status, owner, reviewers, audience, classification, date) fills those
and sets the file's title, subject and author properties. A document with no
`## Document history` table (an older one with `## Document control`, say) has no history page
and keeps that section in the body.

Links to a sibling `x.md` become `x.html` when that page has been built, else plain text, so the
file carries no link that is dead once it is mailed. lint-doc.py runs first; with --strict a lint
error, a mermaid fence with no aws-svg replacement or an unresolved link fails the build.

A DOCX is NOT pixel-identical to the HTML and cannot be: there is no CSS, no rounded panel, no
gradient, and Word breaks pages where its own text engine decides. What carries across is the
design language - Squid Ink headings on white, Smile Orange rules, Squid Ink table headers,
the same content in the same order. The colours are the named constants below.

Requires: pandoc on the PATH, Node.js 22.18+ with the plugin's node packages, and python-docx
in this skill's vendor/python folder - all installed by this skill's setup.sh.
"""

import argparse
import importlib.util
import os
import pathlib
import re
import shutil
import subprocess  # nosec B404 - fixed argument lists of known tools, never a shell
import sys
import tempfile
import urllib.parse

HERE = pathlib.Path(__file__).resolve().parent

# python-docx ships in this skill's vendor/python folder so nothing is installed at run time.
sys.path.insert(0, str(HERE.parent / "vendor" / "python"))
try:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Inches, Mm, Pt, RGBColor
    from docx.text.paragraph import Paragraph
except ImportError as e:
    raise SystemExit(f"error: python-docx is not importable ({e}). Run bash <this skill>/setup.sh "
                     f"with this same python3 ({sys.executable}); lxml is built per Python "
                     "version, so a vendor folder installed by another python3 will not load.") from None


def _load_html_builder():
    """build-html.py owns the markdown rules (markers, standfirst, eyebrow); reuse them."""
    sys.dont_write_bytecode = True   # no __pycache__ left in the skill folder
    spec = importlib.util.spec_from_file_location("build_html", HERE / "build-html.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


HTML = _load_html_builder()

# Palette, AWS brand by default (the values the HTML build uses). For a blue house style set
# HEADING_COLOUR and TABLE_HEAD_FILL to "1F3864" and RULE_COLOUR to "2E74B5", for example.
HEADING_COLOUR = "232F3E"    # Squid Ink: titles and headings
TABLE_HEAD_FILL = "232F3E"   # Squid Ink: table header rows
TABLE_HEAD_TEXT = "FFFFFF"
RULE_COLOUR = "FF9900"       # Smile Orange: the rule under section headings and the cover title
INK = "16191F"
INK_SOFT = "414D5C"
INK_FAINT = "5F6B7A"
LINE = "D5DBDB"              # table grid and the header and contents rules
BG_SOFT = "F2F3F3"           # the label column of the cover table
HEADER_SEP = "  |  "         # between prepared for and prepared by in the page header

BODY_FONT = "Arial"  # on both macOS and Windows. "Segoe UI" is Windows-only and
                     # substitutes to a serif elsewhere, which changes the whole document.
SVG2PNG = HERE.parents[1] / "aws-archify" / "scripts" / "svg-to-png.ts"
NODE_MIN = (22, 18)  # runs .ts files directly (type stripping on by default)


def svg_size(path):
    """(width, height) from the root <svg> tag's viewBox, else its width/height; None if neither."""
    m = re.search(r"""<svg\b(?:[^>"']|"[^"]*"|'[^']*')*>""", path.read_text())
    if not m:
        return None
    root = m.group(0)
    vb = re.search(r"""\sviewBox\s*=\s*["']\s*[-\d.]+[\s,]+[-\d.]+[\s,]+([\d.]+)[\s,]+([\d.]+)""", root)
    if vb:
        return float(vb.group(1)), float(vb.group(2))
    w = re.search(r"""\swidth\s*=\s*["']([\d.]+)(?:px)?["']""", root)
    h = re.search(r"""\sheight\s*=\s*["']([\d.]+)(?:px)?["']""", root)
    return (float(w.group(1)), float(h.group(1))) if w and h else None


def check_node():
    if not shutil.which("node"):
        raise SystemExit("error: Node.js 22.18+ is required to rasterise the figures and `node` "
                         "is not on the PATH.")
    v = subprocess.run(["node", "-p", "process.versions.node"], capture_output=True, text=True)  # nosec B603, B607 - fixed argument list of a tool found on PATH
    parts = tuple(int(x) for x in re.findall(r"\d+", v.stdout)[:2])
    if len(parts) < 2 or parts < NODE_MIN:
        raise SystemExit(f"error: Node.js {v.stdout.strip() or '?'} found; 22.18 or later is "
                         "required to run the .ts converter.")
    if not SVG2PNG.is_file():
        raise SystemExit(f"error: SVG converter not found: {SVG2PNG}")


def rasterise(svg, out, scale=3):
    """SVG -> PNG at `scale`x. Word cannot embed the inline SVG the HTML build uses."""
    if not svg.is_file():
        raise SystemExit(f"error: aws-svg not found: {svg}")
    if re.search(r'<svg\b[^>]*\bdata-poster="true"', svg.read_text()):
        raise SystemExit(f"error: {svg.name} is a poster figure (<Diagram poster>), too wide to read in a page; "
                         "embed an overview and detail figures instead and deliver the poster on its own")
    size = svg_size(svg)
    if not size:
        raise SystemExit(f"error: {svg.name} has no viewBox or width/height on its root <svg>, "
                         "so it cannot be sized for Word.")
    out.unlink(missing_ok=True)  # a PNG left over from an earlier run must not mask a failure
    r = subprocess.run(["node", str(SVG2PNG), str(svg), str(out), str(int(size[0] * scale))],  # nosec B603, B607 - fixed argument list of a tool found on PATH
                       capture_output=True, text=True)
    if r.returncode != 0 or not out.is_file() or out.stat().st_size == 0:
        err = r.stderr.strip()
        hint = ""
        if "@napi-rs/canvas" in err or "ERR_MODULE_NOT_FOUND" in err:
            hint = "\n@napi-rs/canvas is not installed: run bash <plugin>/skills/aws-archify/setup.sh."
        raise SystemExit(f"error: could not rasterise {svg} with {SVG2PNG} "
                         f"(exit {r.returncode})\n{err[-800:]}{hint}")
    return out


def escape_placeholders(line):
    """`<Solution>` outside backticks is raw HTML to pandoc's gfm reader, which drops it."""
    parts = re.split(r"(`+[^`]*`+)", line)
    for n in range(0, len(parts), 2):
        parts[n] = re.sub(r"<(?=[A-Za-z/])(?!(?:https?|mailto):)", "&lt;", parts[n])
    return "".join(parts)


def plain(text):
    """Inline markdown to plain text, for the cover band; code spans keep their text as written."""
    parts = re.split(r"(`+[^`]*`+)", text)
    for n in range(0, len(parts), 2):
        t = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", parts[n])
        parts[n] = re.sub(r"\*\*|(?<!\w)\*|\*(?!\w)", "", t)
    for n in range(1, len(parts), 2):
        parts[n] = parts[n].strip("`")
    return "".join(parts)


def gfm_id(text):
    """The id pandoc's gfm reader gives a heading: lower case, punctuation dropped, spaces to -."""
    text = re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"[*`]", "", text).lower()
    return re.sub(r"\s", "-", re.sub(r"[^\w\s-]", "", text))


# ![alt][ref], ![alt][] and ![ref]: images whose target sits in a reference definition elsewhere.
REF_IMAGE_RE = re.compile(r"!\[[^\]]*\](?:\[[^\]]*\]|(?![(\[]))")
URL_SCHEME_RE = re.compile(r"^[A-Za-z][A-Za-z0-9+.-]*:")


def confined_image(target, src_dir):
    """With --confine, pandoc may embed only an inline data: image or a file inside the document's folder:
    no other file on the server and no URL it would fetch."""
    if target.startswith("data:"):
        return
    if URL_SCHEME_RE.match(target):
        raise SystemExit(f"error: image {target} is a URL (--confine)")
    HTML.confined((src_dir / urllib.parse.unquote(target)).resolve(), target)


def rewrite_links(line, src_dir, anchors, problems):
    """Links as build-html resolves them: x.md -> x.html or plain text, #anchor -> pandoc's id."""
    parts = re.split(r"(`+[^`]*`+)", line)

    def link(m):
        bang, label, target = m.group(1), m.group(2), m.group(3)
        if bang:                        # an image: pandoc resolves it from the source folder
            if HTML.CONFINE_DIR is not None:
                confined_image(target, src_dir)
            return m.group(0)
        href, problem = HTML.resolve_link(target, src_dir, set(anchors))
        if problem and problem not in problems:
            problems.append(problem)
        if href is None:
            return label
        if href.startswith("#") and href[1:] in anchors:
            href = "#" + anchors[href[1:]]
        return f"[{label}]({href})"

    for n in range(0, len(parts), 2):
        if HTML.CONFINE_DIR is not None and REF_IMAGE_RE.search(parts[n]):
            raise SystemExit("error: a reference-style image is not accepted with --confine; write ![alt](path)")
        parts[n] = HTML.LINK_RE.sub(link, parts[n])
    return "".join(parts)


def escape_table_code_pipes(line):
    """A `|` inside a code span in a table row splits the cell in gfm unless written `\\|`."""
    return re.sub(r"(`+)(.+?)\1",
                  lambda m: m.group(1) + re.sub(r"(?<!\\)\|", r"\\|", m.group(2)) + m.group(1), line)


def stage(body, src_dir, work, extra_anchors=()):
    """Rewrite the markdown for pandoc: figures become PNG references, rules go, links resolve.

    extra_anchors are ids an #anchor may name that are not headings here (the history page's
    bookmark). A stray second `# ` heading becomes `##`: pandoc shifts every heading up a level,
    which would turn an H1 into a plain paragraph.

    Returns the staged text, the number of aws-svg figures (so the caller can check every one
    reached the document) and the problems found (unresolved links, mermaid with no figure).
    A captioned figure's alt text is "Figure N. Caption", which style_figures() turns into a
    caption paragraph.
    """
    lines = body.split("\n")
    out, figures, captioned, in_fence, i, problems = [], 0, 0, False, 0, []
    anchors = {}
    for line in lines:                  # build-html's heading ids -> the ids pandoc will write
        hm = HTML.HEADING_RE.match(line.strip())
        if hm:
            anchors.setdefault(HTML.slug(hm.group(2).strip()), gfm_id(hm.group(2).strip()))
    for a in extra_anchors:
        anchors.setdefault(a, a)

    def marker_follows(j):
        """An aws-svg marker at j or j+1 (one blank line allowed), as build-html accepts."""
        for k in (j, j + 1):
            if k >= len(lines):
                return False
            if HTML.marker(lines[k])[0]:
                return True
            if lines[k].strip():
                return False
        return False

    while i < len(lines):
        line, s = lines[i], lines[i].strip()
        if not in_fence and re.match(r"```mermaid\b", s):
            j = i + 1
            while j < len(lines) and not lines[j].strip().startswith("```"):
                j += 1
            if not marker_follows(j + 1):
                msg = ("a mermaid figure has no Word rendering, so its source is shown as code; "
                       "draw it with aws-archify as an AWS-icon SVG to include it here.")
                print(f"  WARNING: {msg}")
                problems.append(msg)
                out += ["```"] + lines[i + 1:j] + ["```"]
            i = j + 1
            continue
        if s.startswith("```"):
            in_fence = not in_fence
        if in_fence or s.startswith("```"):
            out.append(line); i += 1; continue
        kind, path, caption = HTML.marker(s)       # exits on a leftover archify marker
        if kind:
            figures += 1
            png = rasterise(HTML.confined((src_dir / path).resolve(), path), work / f"fig-{figures:02d}.png")
            alt = ""
            if caption:
                captioned += 1
                alt = re.sub(r"[\[\]]", "", plain(caption)).replace("<", "&lt;")
                alt = f"Figure {captioned}. {alt}"
            out += ["", f"![{alt}]({png.name})", ""]  # a name with no spaces: always a valid link
        elif HTML.RULE_RE.match(s):
            out.append("")   # heavy black bars in Word; the heading rules already separate sections
        else:
            if s.startswith("|"):
                line = escape_table_code_pipes(line)
            elif re.match(r"#[ \t]", s):
                line = "#" + s
            out.append(escape_placeholders(rewrite_links(line, src_dir, anchors, problems)))
        i += 1
    return "\n".join(out), figures, problems


PAGE_WIDTH_MM, PAGE_HEIGHT_MM = 210, 297       # A4
TEXT_WIDTH_TWIPS = 11906 - 2 * 1440             # A4 width less the 1in margins set in main()
MONO_FONT = "Consolas"
HEADING_PT = {"1": 14, "2": 13, "3": 12}        # any deeper level is 11
COVER_TOP_PT = 150                              # space above the cover title


def _set(parent, tag, **attrs):
    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(qn(f"w:{k}"), str(v))
    parent.append(el)
    return el


TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
               "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders",
               "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription",
               "tblPrChange"]


def _replace_tblpr(pr, el):
    """Insert into w:tblPr in schema order, replacing any element pandoc already emitted."""
    tag = el.tag.split("}")[1]
    for existing in pr.findall(qn(f"w:{tag}")):
        pr.remove(existing)
    idx = TBLPR_ORDER.index(tag)
    for child in pr:
        ctag = child.tag.split("}")[1]
        if ctag in TBLPR_ORDER and TBLPR_ORDER.index(ctag) > idx:
            child.addprevious(el)
            return
    pr.append(el)


def _tblpr(pr, tag, **attrs):
    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(qn(f"w:{k}"), str(v))
    _replace_tblpr(pr, el)
    return el


def fit_table(table, weights=None):
    """Make a table span the full text column and stop Word re-flowing it.

    Pandoc emits tables 5.5in wide regardless of the page, which reads as a table that does not
    line up with the paragraphs around it. A fixed layout also stops Word widening one column to
    fit a long unbroken token.
    """
    pr = table._tbl.tblPr
    _tblpr(pr, "w:tblW", type="dxa", w=TEXT_WIDTH_TWIPS)
    _tblpr(pr, "w:tblLayout", type="fixed")

    mar = OxmlElement("w:tblCellMar")
    for edge, v in (("top", 70), ("bottom", 70), ("left", 110), ("right", 110)):
        _set(mar, f"w:{edge}", w=v, type="dxa")
    _replace_tblpr(pr, mar)

    grid = table._tbl.find(qn("w:tblGrid"))
    cols = grid.findall(qn("w:gridCol"))
    n = len(cols)
    # A two-column table is always label -> description here, and an even split wastes the
    # description column. Anything wider is left proportional.
    if weights is None:
        weights = [0.32, 0.68] if n == 2 else [1 / n] * n
    for col, wt in zip(cols, weights):
        col.set(qn("w:w"), str(int(TEXT_WIDTH_TWIPS * wt)))

    for row in table.rows:
        _insert_trpr(row._tr.get_or_add_trPr(), OxmlElement("w:cantSplit"))
        for cell, wt in zip(row.cells, weights):
            _set(cell._tc.get_or_add_tcPr(), "w:tcW", w=int(TEXT_WIDTH_TWIPS * wt), type="dxa")


TRPR_ORDER = ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter", "cantSplit",
              "trHeight", "tblHeader", "tblCellSpacing", "jc", "hidden", "ins", "del", "trPrChange"]


def _insert_trpr(pr, el):
    """Insert into w:trPr in schema order (pandoc already writes tblHeader on header rows)."""
    tag = el.tag.split("}")[1]
    if pr.find(qn(f"w:{tag}")) is not None:
        return
    idx = TRPR_ORDER.index(tag)
    for child in pr:
        ctag = child.tag.split("}")[1]
        if ctag in TRPR_ORDER and TRPR_ORDER.index(ctag) > idx:
            child.addprevious(el)
            return
    pr.append(el)


def keep_with_next(paragraph):
    """A heading should never be the last thing on a page."""
    el = OxmlElement("w:keepNext")
    el.set(qn("w:val"), "1")
    _insert_ordered(paragraph._p.get_or_add_pPr(), el)


def centre_images(doc):
    for shp in doc.inline_shapes:
        p = shp._inline.getparent().getparent().getparent()
        pr = p.get_or_add_pPr()
        for existing in pr.findall(qn("w:jc")):
            pr.remove(existing)
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), "center")
        _insert_ordered(pr, jc)


def shade(el, fill):
    """Solid fill on a paragraph or table cell."""
    is_para = el.tag.endswith("}p")
    pr = el.get_or_add_pPr() if is_para else el.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), fill)
    if is_para:
        _insert_ordered(pr, sh)
    else:
        pr.append(sh)


# w:pPr children have a fixed schema order; appending puts them last, which is invalid and can
# make Word drop the entire property set. Anything inserted raw goes through _insert_ordered.
PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
             "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
             "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
             "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
             "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
             "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
             "sectPr", "pPrChange"]


def _insert_ordered(pr, el):
    tag = el.tag.split("}")[1]
    idx = PPR_ORDER.index(tag)
    for child in pr:
        ctag = child.tag.split("}")[1]
        if ctag in PPR_ORDER and PPR_ORDER.index(ctag) > idx:
            child.addprevious(el)
            return
    pr.append(el)


def bottom_rule(paragraph, colour, size=12):
    """The orange rule under each section heading, mirroring the HTML."""
    pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(size))
    b.set(qn("w:space"), "4")
    b.set(qn("w:color"), colour)
    borders.append(b)
    _insert_ordered(pr, borders)


def cell_text(cell, colour=None, bold=None, size=None):
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.name = BODY_FONT
            if colour:
                r.font.color.rgb = RGBColor.from_string(colour)
            if bold is not None:
                r.font.bold = bold
            if size:
                r.font.size = Pt(size)


def table_borders(table, colour=LINE, size=6):
    """Explicit borders rather than the built-in `Table Grid` style: pandoc's output does not
    define that style, and relying on a named style makes the result depend on the reader's
    template."""
    pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(size))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), colour)
        borders.append(e)
    _replace_tblpr(pr, borders)


def style_tables(doc):
    """Light grid, a dark header row in white bold 10 pt that repeats on each page, 10 pt body."""
    for t in doc.tables:
        table_borders(t)
        fit_table(t)
        for i, row in enumerate(t.rows):
            for cell in row.cells:
                if i == 0:
                    shade(cell._tc, TABLE_HEAD_FILL)
                    cell_text(cell, colour=TABLE_HEAD_TEXT, bold=True, size=10)
                else:
                    cell_text(cell, colour=INK, size=10)
        if t.rows:
            hdr = OxmlElement("w:tblHeader")
            hdr.set(qn("w:val"), "1")
            _insert_trpr(t.rows[0]._tr.get_or_add_trPr(), hdr)


def set_styles(doc):
    """Normal and Heading 1-4 in the house font and palette, so text typed later matches.

    Pandoc's reference styles name theme fonts and colours, which win over a plain font name,
    so those attributes are removed first.
    """
    specs = [("Normal", 10.5, False, INK)]
    specs += [(f"Heading {n}", HEADING_PT.get(str(n), 11), True, HEADING_COLOUR) for n in range(1, 5)]
    for name, size, bold, colour in specs:
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        rpr = st.element.get_or_add_rPr()
        for tag in ("w:rFonts", "w:color"):
            for el in rpr.findall(qn(tag)):
                for attr in list(el.attrib):
                    if "theme" in attr.lower():
                        del el.attrib[attr]
        st.font.name, st.font.size, st.font.bold = BODY_FONT, Pt(size), bold
        st.font.italic = False
        st.font.color.rgb = RGBColor.from_string(colour)


def style_headings(doc):
    for p in doc.paragraphs:
        st = (p.style.name or "")
        if st == "Title":
            continue
        if st.startswith("Heading"):
            level = st[-1]          # `##` is Heading 1: pandoc runs with a heading shift of -1
            for r in p.runs:
                r.font.name = BODY_FONT
                r.font.color.rgb = RGBColor.from_string(HEADING_COLOUR)
                r.font.bold = True
                r.font.size = Pt(HEADING_PT.get(level, 11))
            p.paragraph_format.space_before = Pt({"1": 20, "2": 14}.get(level, 10))
            p.paragraph_format.space_after = Pt(6)
            keep_with_next(p)
            if level == "1":
                bottom_rule(p, RULE_COLOUR)
        else:
            p.paragraph_format.space_after = Pt(7)
            p.paragraph_format.line_spacing = 1.15
            for r in p.runs:
                mono = r.style is not None and r.style.name == "Verbatim Char"
                r.font.name = MONO_FONT if mono else BODY_FONT
                if r.font.color.rgb is None:
                    r.font.color.rgb = RGBColor.from_string(INK)
                r.font.size = Pt(9.5) if mono else (r.font.size or Pt(10.5))


COVER_ROWS = (("version", "Version"), ("status", "Status"), ("date", "Date"),
              ("classification", "Classification"), ("owner", "Owner"), ("reviewers", "Reviewers"),
              ("audience", "Audience"))


def new_para(doc, text="", size=10.5, colour=INK, bold=False, centre=False, before=0, after=0):
    """A paragraph appended to the body with one run in the house font; the caller moves it."""
    p = doc.add_paragraph()
    if text:
        r = p.add_run(text)
        r.font.name, r.font.size, r.font.bold = BODY_FONT, Pt(size), bold
        r.font.color.rgb = RGBColor.from_string(colour)
    if centre:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    return p


def page_title(doc, text):
    """A front-matter page title that starts a new page and looks like a Heading 1, but is not
    one, so it stays out of the table of contents."""
    p = new_para(doc, text, HEADING_PT["1"], HEADING_COLOUR, bold=True, after=12)
    p.paragraph_format.page_break_before = True
    bottom_rule(p, RULE_COLOUR)
    return p


def bookmark(doc, paragraph, name):
    """A Word bookmark around a paragraph, so an #anchor link can still land on it."""
    ids = [int(b.get(qn("w:id"))) for b in doc.element.body.iter(qn("w:bookmarkStart"))
           if (b.get(qn("w:id")) or "").isdigit()]
    bid = str(max(ids, default=0) + 1)
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), bid)
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bid)
    paragraph._p.get_or_add_pPr().addnext(start)
    paragraph._p.append(end)


def cover_table(doc, rows):
    """The cover's two-column details table: shaded bold labels, light grid, centred."""
    tbl = doc.add_table(rows=len(rows), cols=2)
    width = int(TEXT_WIDTH_TWIPS * 0.8)
    widths = (int(width * 0.36), width - int(width * 0.36))
    pr = tbl._tbl.tblPr
    _tblpr(pr, "w:tblW", type="dxa", w=width)
    _tblpr(pr, "w:jc", val="center")
    _tblpr(pr, "w:tblLayout", type="fixed")
    mar = OxmlElement("w:tblCellMar")
    for edge, v in (("top", 80), ("bottom", 80), ("left", 140), ("right", 140)):
        _set(mar, f"w:{edge}", w=v, type="dxa")
    _replace_tblpr(pr, mar)
    table_borders(tbl)
    for col, w in zip(tbl._tbl.find(qn("w:tblGrid")).findall(qn("w:gridCol")), widths):
        col.set(qn("w:w"), str(w))
    for row, (label, value) in zip(tbl.rows, rows):
        _insert_trpr(row._tr.get_or_add_trPr(), OxmlElement("w:cantSplit"))
        for cell, w, text in zip(row.cells, widths, (label, value)):
            _set(cell._tc.get_or_add_tcPr(), "w:tcW", w=w, type="dxa")
            cell.text = text
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        shade(row.cells[0]._tc, BG_SOFT)
        cell_text(row.cells[0], colour=INK, bold=True, size=10)
        cell_text(row.cells[1], colour=INK, size=10)
    return tbl


def toc_paragraph(doc):
    """A TOC field (heading levels 1-2, hyperlinked), then a page break to start the body.

    python-docx cannot lay out a table of contents, so the field carries a one-line cached
    result and settings.xml asks Word to update fields on open, which fills it in. The break
    sits after the field's end, so updating the field keeps it.
    """
    toc = doc.add_paragraph()
    shown = add_field(toc, 'TOC \\o "1-2" \\h \\z \\u',
                      "Open in Word and choose Yes to update fields, or right-click here and "
                      "choose Update Field, to build the table of contents.")
    shown.font.name, shown.font.size = BODY_FONT, Pt(9.5)
    shown.font.color.rgb = RGBColor.from_string(INK_SOFT)
    toc.add_run().add_break(WD_BREAK.PAGE)
    update_fields_on_open(doc)
    return toc


def place_first(doc, elements):
    """Move these body elements, in order, to the very top of the document."""
    body = doc.element.body
    for n, el in enumerate(elements):
        body.remove(el)
        body.insert(n, el)


def front_pages(doc, title, subtitle, standfirst, meta, history, anchor=""):
    """Cover page, Document history page (when there is one) and Table of Contents page.

    history says the source had a `## Document history` table; main() put it first in the body,
    so it is the document's first table here. anchor, when set, bookmarks the history page.
    """
    hist_tbl = doc.tables[0]._tbl if history else None
    seq = [new_para(doc, title, 25, HEADING_COLOUR, bold=True, centre=True,
                    before=COVER_TOP_PT, after=10)]
    if subtitle:
        seq.append(new_para(doc, subtitle, 16, INK_SOFT, centre=True, after=10))
    bottom_rule(seq[-1], RULE_COLOUR)
    if standfirst:
        lead = new_para(doc, standfirst, 11, INK_SOFT, centre=True, before=20)
        lead.paragraph_format.left_indent = lead.paragraph_format.right_indent = Inches(0.5)
        seq.append(lead)
    seq.append(new_para(doc, before=40))
    by, to = HTML.prepared(meta)
    rows = [("Prepared by", plain(by))] + ([("Prepared for", plain(to))] if to else [])
    rows += [(label, plain(meta[k])) for k, label in COVER_ROWS if meta.get(k)]
    els = [p._p for p in seq] + [cover_table(doc, rows)._tbl]
    if history:
        h = page_title(doc, "Document history")
        if anchor:
            bookmark(doc, h, anchor)
        els += [h._p, hist_tbl]
    els += [page_title(doc, "Table of Contents")._p, toc_paragraph(doc)._p]
    place_first(doc, els)


def title_block(doc, title, subtitle, standfirst, meta):
    """The one-page layout's compact header: title, subtitle, one metadata line, then the lead."""
    seq = [new_para(doc, title, 20, HEADING_COLOUR, bold=True, after=4)]
    if subtitle:
        seq.append(new_para(doc, subtitle, 12.5, INK_SOFT, after=4))
    line = plain(HTML.meta_text(meta))
    if line:
        seq.append(new_para(doc, line, 8.5, INK_FAINT))
    bottom_rule(seq[-1], RULE_COLOUR)
    seq[-1].paragraph_format.space_after = Pt(12)
    if standfirst:
        seq.append(new_para(doc, standfirst, 10.5, INK, after=8))
    place_first(doc, [p._p for p in seq])


def size_images(doc, max_in=TEXT_WIDTH_TWIPS / 1440):
    for shp in doc.inline_shapes:
        w, h = shp.width, shp.height
        if w > Inches(max_in):
            shp.height = int(h * Inches(max_in) / w)
            shp.width = Inches(max_in)


CAPTION_RE = re.compile(r"^Figure \d+\. ")


def style_figures(doc):
    """A "Figure N. Caption" paragraph under each captioned image; returns how many.

    stage() put the caption in the image's alt text (docPr descr), which is where Word reads
    alt text from, so it stays there and the visible caption is added beneath.
    """
    n = 0
    for shp in doc.inline_shapes:
        descr = shp._inline.docPr.get("descr") or ""
        if not CAPTION_RE.match(descr):
            continue
        n += 1
        img_p = shp._inline.getparent().getparent().getparent()
        keep_with_next(Paragraph(img_p, None))
        cap_el = OxmlElement("w:p")
        img_p.addnext(cap_el)
        cap = Paragraph(cap_el, img_p.getparent())
        num, text = descr.split(". ", 1)
        for t, bold in ((num + ". ", True), (text, False)):
            r = cap.add_run(t)
            r.font.name, r.font.size, r.font.bold = BODY_FONT, Pt(9), bold
            r.font.color.rgb = RGBColor.from_string(INK if bold else INK_SOFT)
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), "center")
        _insert_ordered(cap_el.get_or_add_pPr(), jc)
        cap.paragraph_format.space_before = Pt(4)
        cap.paragraph_format.space_after = Pt(12)
    return n


def add_field(paragraph, instr, placeholder):
    """A complex field (begin, instruction, separate, cached result, end) in one paragraph."""
    def fld(kind):
        r = paragraph.add_run()
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        r._r.append(el)
        return r
    fld("begin")
    r = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r._r.append(it)
    fld("separate")
    shown = paragraph.add_run(placeholder)
    fld("end")
    return shown


# w:settings children that must come after w:updateFields in the schema sequence
SETTINGS_AFTER_UPDATE = {"hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids",
                         "mathPr", "attachedSchema", "themeFontLang", "clrSchemeMapping",
                         "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures", "forceUpgrade",
                         "captions", "readModeInkLockDown", "smartTagType", "schemaLibrary",
                         "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator"}


def update_fields_on_open(doc):
    settings = doc.settings.element
    for old in settings.findall(qn("w:updateFields")):
        settings.remove(old)
    el = OxmlElement("w:updateFields")
    el.set(qn("w:val"), "true")
    for child in settings:
        if child.tag.split("}")[1] in SETTINGS_AFTER_UPDATE:
            child.addprevious(el)
            return
    settings.append(el)


def _hf_paragraph(part, align):
    part.is_linked_to_previous = False
    p = part.paragraphs[0] if part.paragraphs else part.add_paragraph()
    p.text = ""
    p.alignment = align
    return p


def page_header(sec, meta, cover):
    """Right-aligned header: prepared for and prepared by. The cover page has none."""
    by, to = HTML.prepared(meta)
    p = _hf_paragraph(sec.header, WD_ALIGN_PARAGRAPH.RIGHT)
    r = p.add_run(f"{plain(to)}{HEADER_SEP}{plain(by)}" if to else plain(by))
    r.font.name, r.font.size = BODY_FONT, Pt(8.5)
    r.font.color.rgb = RGBColor.from_string(INK_FAINT)
    bottom_rule(p, LINE, size=4)
    if cover:
        sec.different_first_page_header_footer = True
        sec.first_page_header.is_linked_to_previous = False
        sec.first_page_footer.is_linked_to_previous = False


def page_footer(sec, title, meta):
    """Centred footer: "<title> (v<version>) | <classification> | Page N" (a PAGE field)."""
    v = plain(meta.get("version", "")).strip()
    head = plain(title) + (f" ({v if v[:1] in 'vV' else 'v' + v})" if v else "")
    parts = [head] + ([plain(meta["classification"])] if meta.get("classification") else [])
    p = _hf_paragraph(sec.footer, WD_ALIGN_PARAGRAPH.CENTER)
    runs = [p.add_run("  |  ".join(parts) + "  |  Page "), add_field(p, "PAGE", "1")]
    for r in runs:
        r.font.name, r.font.size = BODY_FONT, Pt(8.5)
        r.font.color.rgb = RGBColor.from_string(INK_FAINT)


def core_properties(doc, title, meta):
    cp = doc.core_properties
    cp.title = plain(meta.get("title") or title)
    subject = " / ".join(meta[k] for k in ("status", "classification") if meta.get(k))
    if subject:
        cp.subject = subject
    cp.author = meta.get("owner") or HTML.prepared(meta)[0]


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0],
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("src", help="the markdown design document")
    ap.add_argument("--strict", action="store_true",
                    help="fail (exit 1) on a lint error, a mermaid fence with no aws-svg "
                         "replacement or an unresolved link")
    ap.add_argument("--confine", action="store_true",
                    help="refuse figures and images outside the document's folder (set by a remote MCP server)")
    args = ap.parse_args()
    src = pathlib.Path(args.src).resolve()
    if not src.is_file():
        raise SystemExit(f"error: {src} not found")
    if args.confine:
        HTML.CONFINE_DIR = src.parent
    if not shutil.which("pandoc"):
        raise SystemExit("error: pandoc is required for Word output. Install it with "
                         "`brew install pandoc` (macOS), `sudo apt-get install pandoc` or "
                         "`sudo dnf install pandoc` (Linux), or see https://pandoc.org/installing.html")

    if HTML.lint_doc(src) and args.strict:
        raise SystemExit("error: lint-doc.py found errors; fix them, or build without --strict")

    meta, md = HTML.front_matter(src.read_text().replace("\r\n", "\n"))
    tm = re.search(r"^# +(.+)$", md, re.M)
    if not tm and not meta.get("title"):
        raise SystemExit(f"error: {src.name} has no `# Title` line (or front matter title); "
                         "the Word cover needs one.")
    title = tm.group(1).strip() if tm else meta["title"]
    one_page = HTML.doc_kind(meta) in HTML.ONE_PAGE_KINDS
    # The same paragraph build-html.py lifts into its header, removed from the body the same way.
    body, standfirst = HTML.split_lead(md)
    body = re.sub(r"^# +.+$", "", body, count=1, flags=re.M)      # title comes from the cover
    history, hid = [], ""
    if not one_page:
        # the history table goes first, so it is the first table pandoc writes (front_pages)
        body, history, hid = HTML.split_history(body)
        if history:
            body = "\n".join(history) + "\n\n" + body

    out = src.parent / (src.stem.lower() + ".docx")
    with tempfile.TemporaryDirectory() as tmp:
        work = pathlib.Path(tmp)
        if re.search(r"^\s*<!--\s*aws-svg:", body, re.M):
            check_node()
        staged_md, figures, problems = stage(body, src.parent, work, [hid] if hid else [])
        for p in problems:
            if not p.startswith("a mermaid"):
                print(f"  WARNING: {p}")
        if problems and args.strict:
            raise SystemExit(f"error: {len(problems)} problem(s) above (--strict)")
        (work / "staged.md").write_text(staged_md)
        # Fixed pandoc argument list, no shell.
        r = subprocess.run(  # nosec B603, B607 - fixed argument list of a tool found on PATH
            ["pandoc", "staged.md", "--from", "gfm", "--to", "docx", "--shift-heading-level-by=-1",
             "--resource-path", os.pathsep.join([str(work), str(src.parent)]), "-o", str(out)],
            cwd=work, capture_output=True, text=True)
        if r.returncode != 0:
            raise SystemExit(f"error: pandoc failed (exit {r.returncode})\n{r.stderr.strip()[-800:]}")

    doc = Document(str(out))
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(PAGE_WIDTH_MM), Mm(PAGE_HEIGHT_MM)
    sec.left_margin = sec.right_margin = Inches(1.0)
    sec.top_margin = sec.bottom_margin = Inches(0.9)

    set_styles(doc)
    style_headings(doc)
    style_tables(doc)   # before front_pages(): the cover table has its own style
    size_images(doc)
    centre_images(doc)
    captions = style_figures(doc)
    subtitle = plain(meta.get("subtitle", ""))
    if one_page:
        title_block(doc, plain(title), subtitle, plain(standfirst), meta)
    else:
        front_pages(doc, plain(title), subtitle, plain(standfirst), meta, bool(history), hid)
    page_header(sec, meta, cover=not one_page)
    page_footer(sec, meta.get("title") or title, meta)
    core_properties(doc, title, meta)
    doc.save(str(out))

    images = len(doc.inline_shapes)
    kb = out.stat().st_size // 1024
    print(f"wrote {out.name}  {kb} KB  "
          f"{len(doc.tables)} tables  {images} images  {figures} aws-svg  {captions} captioned  "
          f"{sum(1 for p in doc.paragraphs if (p.style.name or '').startswith('Heading'))} headings  "
          f"{len(problems)} problems")
    if images < figures:
        raise SystemExit(f"error: {figures} aws-svg figure(s) in the source but only {images} "
                         "image(s) reached the .docx")


if __name__ == "__main__":
    main()
